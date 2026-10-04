"""Generates mock import files into src/lib/__tests__/fixtures (run: python3 scripts/gen-fixtures.py)."""
import datetime as dt
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font

out = Path(__file__).resolve().parent.parent / "src/lib/__tests__/fixtures"
out.mkdir(parents=True, exist_ok=True)
F = Font(name="Arial")

def sheet(ws, rows):
    for r in rows:
        ws.append(r)
    for row in ws.iter_rows():
        for c in row:
            c.font = F

# ---- workers.xlsx (cols: id, rank, last, first, -, release, -,-,-, phone, notes, eligibility)
wb = Workbook(); ws = wb.active; ws.title = "עובדים"
H = ["מספר אישי", "דרגה", "שם משפחה", "שם פרטי", "", "תאריך שחרור", "", "", "", "מספר טלפון נייד", "הערות", "כשירות"]
sheet(ws, [
    H,
    [1001, "סמל", "כהן", "דני", None, dt.date(2026, 12, 31), None, None, None, 501234567, "יום לימודים ג'", "כשיר"],
    [1002, "סמל", "לוי", "רונית", None, None, None, None, None, "052-765-4321", None, "כשיר"],
    [1003, "רב סמל", "ישראלי", "יוסי", None, None, None, None, None, None, "פציעה בברך", "לא כשיר"],
    [1004, "סמל", None, "חסר", None, None, None, None, None, None, None, "כשיר"],      # missing last name -> error
    [None, "סמל", "ריק", "מזהה", None, None, None, None, None, None, None, "כשיר"],    # no id -> skipped
    ["A-77", "סגן", "בן דוד", "אורי", None, 0, None, None, None, "+972501112222", None, None],  # release 0 -> null
])
wb.save(out / "workers.xlsx")

wb = Workbook(); wb.active.title = "ריק"; wb.save(out / "workers-empty.xlsx")

wb = Workbook(); ws = wb.active; ws.title = "noheader"
sheet(ws, [[2001, "סמל", "אבי", "נועם"]]); wb.save(out / "workers-noheader.xlsx")

# ---- shift-dates.xlsx
wb = Workbook(); ws = wb.active; ws.title = "בוקר"
sheet(ws, [["תאריך"], [dt.date(2026, 1, 4)], [dt.date(2026, 1, 8)], ["05/01/2026"], ["לא תאריך"], [None], ["9-10/1/2026"], ["15.1-16.1.2026"]])
ws2 = wb.create_sheet("ערב")
sheet(ws2, [[dt.date(2026, 2, 1)], [dt.date(2026, 2, 5)]])
wb.create_sheet("ריק")  # no data -> skipped
ws4 = wb.create_sheet("זבל"); sheet(ws4, [["abc"], ["xyz"]])
wb.save(out / "shift-dates.xlsx")

# ---- form responses
wb = Workbook(); ws = wb.active
sheet(ws, [
    ["חותמת זמן", "שם", "ענף", "אילוצים"],
    ["01/10/2026 10:00", "דני כהן", "א", "לא יכול 5.1-7.1"],
    ["01/10/2026 10:05", "  לוי רונית ", "ב", ""],
    ["01/10/2026 10:06", "", "ב", "שורה בלי שם"],
    ["01/10/2026 10:07", "יוסי ישראלי", "ג", "שורה1\nשורה2"],
])
wb.save(out / "form-responses.xlsx")

# CSV fixtures (with BOM, CRLF, tabs, comments)
(out / "workers.csv").write_text(
    "﻿worker_id,name,rank_id,is_exempt,exemption_reason,receives\r\n"
    "# comment\r\n"
    "1001,דני כהן,סמל,0\r\n"
    "1002,רונית לוי,סמל,1,פטור רפואי,0\r\n"
    "1003\tיוסי\tרב סמל\r\n"
    "1004,,סמל\r\n"
    "1005,אורי,לא-קיים\r\n", encoding="utf-8")
(out / "shift-dates.csv").write_text(
    "﻿תאריך,סוג\r\n2026-01-08,בוקר\r\n05/01/2026\r\n# x\r\nbad\r\n07.01.26,ערב\r\n", encoding="utf-8")
(out / "form-responses.csv").write_text(
    "﻿חותמת זמן,שם,ענף,אילוצים\r\n01/10/2026,דני כהן,א,\"לא יכול 5.1\"\r\n01/10/2026,רונית לוי,ב,\r\n", encoding="utf-8")
print("ok")
